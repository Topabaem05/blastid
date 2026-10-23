"""YAML 출력 템플릿과 헤더가 들어 있는 xlsx 템플릿."""
from __future__ import annotations

from dataclasses import dataclass
from dataclasses import field as dc_field
from pathlib import Path

import yaml

from blastid.fields import FIELD_HELP

DATA_DIR = Path(__file__).resolve().parent / "data"


@dataclass
class Column:
    header: str
    field: str | None = None
    number_format: str | None = None
    width: float | None = None
    role: str = "result"
    aliases: list[str] = dc_field(default_factory=list)
    wrap: bool = False

    def names(self) -> list[str]:
        out = []
        for n in [self.header, *self.aliases]:
            if n and n not in out:
                out.append(n)
        return out


@dataclass
class Template:
    sheet: str = "Sheet1"
    screenshot_sheet: str = "스크린샷"
    header_row: int = 1
    freeze: str = "D2"
    row_height: float = 30
    header_height: float = 32
    autofilter: bool = True
    bold_field: str = "id"
    empty_field: str = "scientific_name"
    id_field: str = "id"
    sequence_field: str = "sequence"
    comment_field: str = "comment"
    screenshot_field: str = "screenshot"
    style: dict = dc_field(default_factory=dict)
    columns: list[Column] = dc_field(default_factory=list)
    workbook_path: Path | None = None
    source: Path | None = None

    def column_for(self, field_name: str) -> Column | None:
        for c in self.columns:
            if c.field == field_name:
                return c
        return None


def packaged_template_path() -> Path:
    return DATA_DIR / "default.yaml"


def _color(value: str | None, default: str) -> str:
    raw = (value or default).strip()
    return raw[1:] if raw.startswith("#") else raw


def load_template(path: str | Path | None = None) -> Template:
    src = Path(path) if path else packaged_template_path()
    if not src.exists():
        raise SystemExit(f"템플릿 파일이 없습니다: {src}")
    data = yaml.safe_load(src.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise SystemExit(f"템플릿 YAML 형식이 잘못되었습니다: {src}")
    columns = []
    for i, raw in enumerate(data.get("columns") or [], start=1):
        if not isinstance(raw, dict):
            raise SystemExit(f"{src}: columns[{i}] 가 사전이 아닙니다.")
        header = raw.get("header")
        if header is None:
            raise SystemExit(f"{src}: columns[{i}] 에 header 가 없습니다.")
        fld = raw.get("field") or None
        if fld and fld not in FIELD_HELP:
            known = ", ".join(sorted(FIELD_HELP))
            raise SystemExit(f"{src}: 알 수 없는 field '{fld}'. 쓸 수 있는 값: {known}")
        role = raw.get("role") or ("input" if fld in {"no", "id", "sequence"} else "result")
        columns.append(
            Column(
                header=str(header),
                field=fld,
                number_format=raw.get("number_format") or None,
                width=raw.get("width"),
                role=role,
                aliases=[str(a) for a in (raw.get("aliases") or [])],
                wrap=bool(raw.get("wrap", fld in {"lineage", "comment", "scientific_name", "sequence"})),
            )
        )
    if not columns:
        raise SystemExit(f"{src}: columns 가 비어 있습니다.")
    book = data.get("workbook")
    book_path = None
    if book:
        cand = (src.parent / str(book)).resolve()
        if cand.exists():
            book_path = cand
    style = dict(data.get("style") or {})
    return Template(
        sheet=str(data.get("sheet") or "Sheet1"),
        screenshot_sheet=str(data.get("screenshot_sheet") or "스크린샷"),
        header_row=int(data.get("header_row") or 1),
        freeze=str(data.get("freeze") or "D2"),
        row_height=float(data.get("row_height") or 30),
        header_height=float(data.get("header_height") or 32),
        autofilter=bool(data.get("autofilter", True)),
        bold_field=str(data.get("bold_field") or "id"),
        empty_field=str(data.get("empty_field") or "scientific_name"),
        id_field=str(data.get("id_field") or "id"),
        sequence_field=str(data.get("sequence_field") or "sequence"),
        comment_field=str(data.get("comment_field") or "comment"),
        screenshot_field=str(data.get("screenshot_field") or "screenshot"),
        style=style,
        columns=columns,
        workbook_path=book_path,
        source=src,
    )


def build_workbook(template: Template):
    """헤더와 서식만 있는 새 워크북."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = template.sheet[:31] or "Sheet1"
    style = template.style
    header_fill = PatternFill("solid", fgColor=_color(style.get("header_fill"), "1F4E78"))
    font = Font(bold=True, color=_color(style.get("header_font"), "FFFFFF"), size=int(style.get("header_size") or 11))
    thin = Side(style="thin", color=_color(style.get("border"), "BFBFBF"))
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hr = template.header_row
    for i, col in enumerate(template.columns, start=1):
        cell = ws.cell(hr, i, col.header)
        cell.fill = header_fill
        cell.font = font
        cell.alignment = align
        cell.border = border
        if col.width:
            ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.row_dimensions[hr].height = template.header_height
    ws.freeze_panes = template.freeze
    if template.autofilter and template.columns:
        last = get_column_letter(len(template.columns))
        ws.auto_filter.ref = f"A{hr}:{last}{hr}"
    return wb


def init_template(dest: Path) -> tuple[Path, Path]:
    """기본 YAML과 xlsx 템플릿을 dest 에 쓴다."""
    dest.mkdir(parents=True, exist_ok=True)
    raw = packaged_template_path().read_text(encoding="utf-8")
    raw = raw.replace("workbook: default.xlsx", "workbook: template.xlsx")
    yml = dest / "template.yaml"
    xlsx = dest / "template.xlsx"
    yml.write_text(raw, encoding="utf-8")
    tmpl = load_template(yml)
    wb = build_workbook(tmpl)
    wb.save(xlsx)
    return yml, xlsx
