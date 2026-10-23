"""엑셀 열 매칭, 셀 기록, 서식, 스크린샷 시트."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from blastid.fields import INPUT_FIELDS, LINK_FIELDS
from blastid.template import Template, _color, build_workbook

SEQ_RE = re.compile(r"^[ACGTURYKMSWBDHVN]+$", re.I)


def norm_header(value) -> str:
    if value is None:
        return ""
    s = str(value).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", s).casefold()


def clean_seq(value, min_len: int = 50) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", "", str(value)).upper().replace("U", "T")
    if len(s) >= min_len and SEQ_RE.match(s):
        return s
    return None


def parse_rows(spec: str) -> list[int]:
    rows: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            rows.update(range(int(a), int(b) + 1))
        else:
            rows.add(int(part))
    return sorted(rows)


def header_map(ws, header_row: int) -> dict[str, int]:
    found: dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        key = norm_header(ws.cell(header_row, col).value)
        if key and key not in found:
            found[key] = col
    return found


def map_fields(ws, template: Template) -> dict[str, int]:
    found = header_map(ws, template.header_row)
    out: dict[str, int] = {}
    for col in template.columns:
        if not col.field:
            continue
        for name in col.names():
            key = norm_header(name)
            if key and key in found:
                out[col.field] = found[key]
                break
    return out


def guess_sequence_column(ws, header_row: int, min_len: int) -> int | None:
    best = None
    last = min(ws.max_row or header_row, header_row + 30)
    for col in range(1, (ws.max_column or 1) + 1):
        hits = 0
        for row in range(header_row + 1, last + 1):
            if clean_seq(ws.cell(row, col).value, min_len):
                hits += 1
        if hits and (best is None or hits > best[0]):
            best = (hits, col)
    return best[1] if best else None


def guess_id_column(ws, header_row: int, seq_col: int | None) -> int | None:
    for col in range(1, (ws.max_column or 1) + 1):
        if col == seq_col:
            continue
        h = norm_header(ws.cell(header_row, col).value)
        if not h:
            continue
        if h in {"id", "voucher", "sample", "sample id"} or "voucher" in h or h.startswith("id"):
            return col
    return None


def ensure_fields(ws, template: Template, field_cols: dict[str, int], log=print) -> dict[str, int]:
    """템플릿의 result/meta 필드가 없으면 시트 끝에 헤더를 추가한다."""
    found = header_map(ws, template.header_row)
    out = dict(field_cols)
    next_col = (ws.max_column or 0) + 1
    style = template.style
    fill = PatternFill("solid", fgColor=_color(style.get("header_fill"), "1F4E78"))
    font = Font(bold=True, color=_color(style.get("header_font"), "FFFFFF"), size=int(style.get("header_size") or 11))
    thin = Side(style="thin", color=_color(style.get("border"), "BFBFBF"))
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    added = []
    for col in template.columns:
        if not col.field or col.role == "input" or col.field in out:
            continue
        # 같은 헤더가 다른 철자로 이미 있으면 map 이 잡았을 것. 여기까지 왔으면 없음.
        if norm_header(col.header) in found:
            out[col.field] = found[norm_header(col.header)]
            continue
        cell = ws.cell(template.header_row, next_col, col.header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        if col.width:
            ws.column_dimensions[get_column_letter(next_col)].width = col.width
        out[col.field] = next_col
        added.append(col.header)
        next_col += 1
    if added:
        log("기존 시트에 없어 오른쪽에 열을 추가: " + ", ".join(added))
    return out


def pick_sheet(wb, name: str, log=print):
    if name in wb.sheetnames:
        return wb[name]
    if len(wb.sheetnames) == 1:
        log(f"시트 '{name}' 이 없어 '{wb.sheetnames[0]}' 시트를 사용합니다.")
        return wb[wb.sheetnames[0]]
    raise SystemExit(f"시트 '{name}' 이 없습니다. 있는 시트: {', '.join(wb.sheetnames)}")


def new_workbook(template: Template):
    import openpyxl

    if template.workbook_path and template.workbook_path.exists():
        return openpyxl.load_workbook(template.workbook_path)
    return build_workbook(template)


def write_fields(ws, row: int, template: Template, field_cols: dict[str, int], values: dict, *, fill_input: bool):
    link_font = Font(color=_color(template.style.get("link_font"), "0563C1"), underline="single")
    for coldef in template.columns:
        if not coldef.field or coldef.field not in field_cols:
            continue
        if coldef.field == template.screenshot_field:
            continue  # 스크린샷 칸은 place_shot 이 채운다
        if not fill_input and coldef.field in INPUT_FIELDS:
            continue
        value = values.get(coldef.field)
        if value == "":
            value = None
        # 빈 코멘트로 사람이 적어 둔 칸을 지우지 않는다. 지울 때는 clear_review.
        if coldef.field == template.comment_field and not value:
            continue
        cell = ws.cell(row, field_cols[coldef.field])
        cell.value = value
        if coldef.number_format and cell.value is not None:
            cell.number_format = coldef.number_format
        if coldef.field in LINK_FIELDS and cell.value:
            cell.hyperlink = str(cell.value)
            cell.font = link_font
        if coldef.field == template.comment_field and cell.value:
            cell.comment = Comment(str(cell.value), "blastid", width=420, height=160)


def should_clear_comment(cell, clear_old: bool) -> bool:
    """문제가 없어진 행에서 blastid 가 남긴 코멘트만 지운다.

    ``--clear-old-comments`` 이면 그 칸의 기존 코멘트를 지운다.
    """
    if clear_old:
        return cell.value not in (None, "") or cell.comment is not None or cell.hyperlink is not None
    text = "" if cell.value is None else str(cell.value)
    author = ""
    if cell.comment is not None and cell.comment.author:
        author = str(cell.comment.author)
    return text.startswith("[") and "blastid" in author.lower()


def clear_review(ws, row: int, template: Template, field_cols: dict[str, int]):
    for fld in (template.comment_field, template.screenshot_field):
        col = field_cols.get(fld)
        if not col:
            continue
        cell = ws.cell(row, col)
        cell.value = None
        cell.comment = None
        cell.hyperlink = None


def level_from_comment(text) -> str | None:
    if not isinstance(text, str):
        return None
    if text.startswith("[이상]"):
        return "strange"
    if text.startswith("[애매]"):
        return "ambiguous"
    return None


def _img_end_row(img) -> int:
    anchor = img.anchor
    if isinstance(anchor, str):
        r0 = int(re.sub(r"[A-Z]", "", anchor))
    else:
        r0 = anchor._from.row + 1
    return r0 + int((img.height or 600) / 20) + 2


def shot_sheet(wb, title: str):
    if title in wb.sheetnames:
        sh = wb[title]
        nxt = max([sh.max_row + 2] + [_img_end_row(i) + 1 for i in sh._images])
    else:
        sh = wb.create_sheet(title)
        sh["A1"] = "BLAST 결과 목록(상위 15개) 이미지 — 문제 행"
        sh["A1"].font = Font(bold=True, size=13)
        sh.column_dimensions["A"].width = 140
        nxt = 3
    return sh, nxt


def place_shot(wb, ws, row: int, png: Path, caption: str, template: Template, field_cols: dict[str, int]):
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PImg

    sh, at = shot_sheet(wb, template.screenshot_sheet)
    sh.cell(at, 1, caption).font = Font(bold=True)
    img = XLImage(str(png))
    w, h = PImg.open(png).size
    scale = min(1.0, 1100 / w)
    img.width, img.height = int(w * scale), int(h * scale)
    sh.add_image(img, f"A{at + 1}")
    col = field_cols.get(template.screenshot_field)
    if not col:
        return
    label = f"스크린샷 보기 ({row}행)"
    cell = ws.cell(row, col, label)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{template.screenshot_sheet}'!A{at}", display=label)
    cell.font = Font(color=_color(template.style.get("link_font"), "0563C1"), underline="single")


def style_sheet(ws, template: Template, field_cols: dict[str, int], flags: dict[int, str]):
    style = template.style
    header_fill = PatternFill("solid", fgColor=_color(style.get("header_fill"), "1F4E78"))
    header_font = Font(bold=True, color=_color(style.get("header_font"), "FFFFFF"), size=int(style.get("header_size") or 11))
    thin = Side(style="thin", color=_color(style.get("border"), "BFBFBF"))
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    band = PatternFill("solid", fgColor=_color(style.get("band_fill"), "EEF3F8"))
    res = PatternFill("solid", fgColor=_color(style.get("result_fill"), "E2EFDA"))
    yel = PatternFill("solid", fgColor=_color(style.get("ambiguous_fill"), "FFF2CC"))
    red = PatternFill("solid", fgColor=_color(style.get("strange_fill"), "F8CBAD"))
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hr = template.header_row
    last_col = max([ws.max_column or 1, *field_cols.values()])
    result_cols = {
        field_cols[c.field]
        for c in template.columns
        if c.field and c.role == "result" and c.field in field_cols
    }
    wrap_cols = {
        field_cols[c.field]
        for c in template.columns
        if c.field and c.wrap and c.field in field_cols
    }
    seq_col = field_cols.get(template.sequence_field)
    name_col = field_cols.get(template.empty_field)
    for c in range(1, last_col + 1):
        cell = ws.cell(hr, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[hr].height = template.header_height
    widths = {field_cols[c.field]: c.width for c in template.columns if c.field and c.field in field_cols and c.width}
    # 필드가 없는 입력 헤더(Collection No. 등)도 기본 템플릿 폭을 맞춘다.
    found = header_map(ws, hr)
    for c in template.columns:
        if c.width and norm_header(c.header) in found:
            widths.setdefault(found[norm_header(c.header)], c.width)
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    last = ws.max_row or hr
    for r in range(hr + 1, last + 1):
        if all(ws.cell(r, c).value is None for c in range(1, last_col + 1)):
            continue
        lvl = flags.get(r)
        named = name_col is not None and ws.cell(r, name_col).value not in (None, "")
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=(c in wrap_cols))
            if c == seq_col:
                continue
            if c in result_cols and lvl == "strange":
                cell.fill = red
            elif c in result_cols and lvl == "ambiguous":
                cell.fill = yel
            elif c in result_cols and named:
                cell.fill = res
            elif r % 2 == 0:
                cell.fill = band
        ws.row_dimensions[r].height = template.row_height
    ws.freeze_panes = template.freeze
    if template.autofilter:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(last_col)}{last}"


def column_letter(template: Template, field_cols: dict[str, int], field_name: str) -> str | None:
    col = field_cols.get(field_name)
    if not col:
        return None
    return get_column_letter(col)
