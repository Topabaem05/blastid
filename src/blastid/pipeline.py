"""서열을 모아 BLAST하고 템플릿 엑셀에 쓴다."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import re
import time
from pathlib import Path

from blastid.assemble import assemble_files, write_qc
from blastid.fields import build_field_values
from blastid.localblast import (
    LocalBlastUnavailable,
    local_note,
    run_local,
    source_label,
)
from blastid.ncbi import DEFAULT_MAX_WAIT, NCBI, WEB_LINK, NcbiWaitTimeout, blast_reports, check_params, rid_url
from blastid.policy import DISCLAIMER_KO, UsageLog, plan_submission, remote_batch_budget, sleep_until
from blastid.policy import ET as ET_ZONE
from blastid.scoring import assess, fungi_lineage, parse_xml_report, web_values
from blastid.template import Template, load_template
from blastid.workbook import (
    clean_seq,
    clear_review,
    ensure_fields,
    guess_id_column,
    guess_sequence_column,
    level_from_comment,
    map_fields,
    new_workbook,
    norm_header,
    parse_rows,
    pick_sheet,
    place_shot,
    should_clear_comment,
    style_sheet,
    write_fields,
)

LOG_FIELDS = [
    "row", "id", "rid", "web_link", "query_index", "qlen", "name", "max", "total",
    "query_cover", "evalue", "per_ident", "acc_len", "accession", "taxid", "lineage",
    "flag", "comment", "verified", "elapsed_s", "merge_mode", "merge_backend",
    "direction", "overlap_len", "overlap_ident", "n_conflicts", "n_ambiguities",
    "trimmed_len", "source", "search_backend",
]


def default_outdir(base) -> Path:
    d = Path(base or Path.cwd() / "blast_runs") / dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    d.mkdir(parents=True, exist_ok=True)
    return d


def make_logger(outdir: Path):
    fh = open(Path(outdir) / "run.log", "a", encoding="utf-8")

    def log(msg):
        line = f"[{dt.datetime.now():%H:%M:%S}] {msg}"
        print(line, flush=True)
        fh.write(line + "\n")
        fh.flush()

    return log


def fasta_label(text, fallback: str) -> str:
    s = re.sub(r"\s+", "_", str(text or "").strip())
    s = re.sub(r"[^\w.\-+]", "_", s, flags=re.UNICODE)
    s = s.strip("_.") or fallback
    return s[:80]


def unique_labels(queries: list[dict]) -> None:
    seen: dict[str, int] = {}
    for q in queries:
        base = q["label"]
        n = seen.get(base, 0) + 1
        seen[base] = n
        if n > 1:
            q["label"] = f"{base}__{n}"


def run_blast(
    ncbi,
    queries,
    batch_size,
    outdir,
    log,
    reuse=None,
    usage: UsageLog | None = None,
    *,
    backend: str = "auto",
    remote_limit: int = 100,
    local_db: str = "ITS_RefSeq_Fungi",
    blastn: str | None = None,
    max_wait: float = DEFAULT_MAX_WAIT,
) -> dict:
    raw = Path(outdir) / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    chunks = [queries[i : i + batch_size] for i in range(0, len(queries), batch_size)]
    if reuse and len(reuse) != len(chunks):
        raise SystemExit(f"--reuse-rids 개수({len(reuse)})가 제출 묶음 수({len(chunks)})와 다릅니다.")
    usage = usage or UsageLog()
    mode = "remote" if reuse else (backend or "auto")
    remote_n = remote_batch_budget(usage.count(), len(chunks), remote_limit, mode)
    if mode == "remote" and not reuse and usage.count() + len(chunks) > remote_limit:
        log(
            f"NCBI: --backend remote 로 24시간 한도({remote_limit})를 넘겨 제출합니다. "
            "NCBI가 느리게 하거나 막을 수 있습니다. IP는 바꾸지 않습니다."
        )
    if remote_n < len(chunks):
        log(
            f"최근 24시간 원격 검색 {usage.count()}건. 한도 {remote_limit}건이라 "
            f"앞 {remote_n}묶음만 NCBI 웹으로 보내고, 나머지는 로컬 BLAST+ ({local_db}) 입니다."
        )
    jobs = {}
    chunk_of = {}
    for ci, ch in enumerate(chunks[:remote_n]):
        fasta = "".join(f">{q['label']}\n{q['seq']}\n" for q in ch)
        if reuse:
            rid, rtoe = reuse[ci], 0
            ncbi._last_poll.pop(rid, None)
            log(f"기존 RID 재사용: {rid} ({len(ch)}개 서열) {WEB_LINK.format(rid=rid)}")
        else:
            rid, rtoe = ncbi.submit(fasta)
            if usage is not None:
                try:
                    usage.add(rid=rid)
                except OSError as e:
                    log(f"  검색 횟수 기록 실패({e})")
            log(f"제출: RID {rid} ({len(ch)}개 서열, 예상 {rtoe}초) {WEB_LINK.format(rid=rid)}")
        jobs[rid] = (time.time(), rtoe)
        chunk_of[rid] = ch
    timed_out: list[str] = []
    try:
        ready = ncbi.wait_all(jobs, max_wait=max_wait) if jobs else {}
    except NcbiWaitTimeout as exc:
        ready = exc.ready
        timed_out = list(exc.rids)
        for rid in timed_out:
            log(str(exc) if rid == timed_out[0] else f"RID {rid} 대기 초과 {WEB_LINK.format(rid=rid)}")
    results = {}
    for rid, ch in chunk_of.items():
        if rid not in ready:
            continue
        payload = ncbi.get_json(rid)
        (raw / f"{rid}.json").write_text(json.dumps(payload), encoding="utf-8")
        reps = blast_reports(payload)
        if len(reps) != len(ch):
            raise RuntimeError(f"RID {rid}: 쿼리 {len(ch)}개인데 결과 {len(reps)}개")
        for qi, (q, rep) in enumerate(zip(ch, reps)):
            report = rep["report"]
            search = report["results"]["search"]
            title = (search.get("query_title") or "").split()
            if not title or title[0] != q["label"].split()[0]:
                raise RuntimeError(
                    f"RID {rid}: 쿼리 순서 불일치 {search.get('query_title')!r} != {q['label']!r}"
                )
            qlen = int(search["query_len"])
            hits = []
            for h in search.get("hits") or []:
                if h.get("hsps"):
                    hits.append(web_values(h, qlen))
            target = report.get("search_target") or {}
            results[q["key"]] = dict(
                rid=rid,
                rid_url=rid_url(rid),
                query_index=qi + 1,
                qlen=qlen,
                seq_len=len(q["seq"]),
                db=target.get("db"),
                backend="remote",
                source=source_label(target.get("db") or "core_nt", local=False, rid=rid),
                params=report.get("params"),
                blast_version=report.get("version"),
                hits=hits,
                elapsed=ready[rid] - jobs[rid][0],
            )
    if timed_out:
        late = [q for rid in timed_out for q in chunk_of.get(rid, [])]
        links = " ".join(f"{rid} {WEB_LINK.format(rid=rid)}" for rid in timed_out)
        reuse_flag = f"--reuse-rids {','.join(timed_out)}"
        if mode == "auto" and late:
            log(f"대기가 길어 이 {len(late)}개 서열은 로컬 BLAST+ ({local_db})로 넘깁니다.")
            try:
                results.update(run_local(late, local_db, Path(outdir), log, blastn=blastn))
            except LocalBlastUnavailable as exc:
                raise LocalBlastUnavailable(
                    f"{exc} NCBI는 아직 WAITING 입니다. {links} 나중에 {reuse_flag}",
                    {**results, **(exc.partial or {})},
                ) from exc
        elif late:
            raise LocalBlastUnavailable(
                f"NCBI가 {max_wait / 60:.0f}분 안에 READY가 되지 않았습니다. {links} "
                f"나중에 {reuse_flag}",
                results,
            )
    if remote_n < len(chunks):
        rest = [q for ch in chunks[remote_n:] for q in ch]
        try:
            local_results = run_local(rest, local_db, Path(outdir), log, blastn=blastn)
        except LocalBlastUnavailable as exc:
            exc.partial = results
            raise
        for item in local_results.values():
            item["elapsed"] = item.get("elapsed") or 0
        results.update(local_results)
        log(f"로컬 BLAST+ ({local_db}) {len(rest)}개 서열 완료")
    return results


def add_taxonomy(ncbi, results, log) -> None:
    tax = ncbi.taxonomy(r["hits"][0]["taxid"] for r in results.values() if r["hits"])
    for r in results.values():
        if not r["hits"]:
            r["lineage_full"] = r["lineage"] = r["tax_name"] = ""
            continue
        t = tax.get(str(r["hits"][0]["taxid"]))
        r["lineage_full"] = t["lineage_full"] if t else ""
        r["lineage"] = fungi_lineage(r["lineage_full"])
        r["tax_name"] = t["name"] if t else ""
        if t and t["name"] != r["hits"][0]["name"]:
            log(
                f"  (참고) RID {r['rid']} query {r['query_index']}: "
                f"BLAST 표시명 '{r['hits'][0]['name']}' / Taxonomy 현재명 '{t['name']}'"
            )


def combine_review(blast_level, blast_comment, merge):
    merge_text = ""
    merge_level = None
    if merge is not None:
        merge_text = merge.warning or ""
        if merge.mode == "fallback":
            merge_level = "ambiguous"
        elif merge.mode == "consensus" and (merge.n_ambiguities or 0) >= 3:
            merge_level = "ambiguous"
            if merge_text and not merge_text.startswith("["):
                merge_text = "[애매] " + merge_text
                merge.warning = merge_text
    level = blast_level or merge_level
    parts = []
    if blast_comment and merge_text:
        parts = [blast_comment, merge_text] if blast_level else [merge_text, blast_comment]
    elif blast_comment:
        parts = [blast_comment]
    elif merge_text:
        parts = [merge_text]
    return level, " ".join(parts)


def make_record(row, sample_id, res, level, comment, merge) -> dict:
    t = res["hits"][0] if res["hits"] else {}
    rec = dict(
        row=row,
        id=sample_id,
        rid=res["rid"],
        web_link=res["rid_url"],
        query_index=res["query_index"],
        qlen=res["qlen"],
        seq_len=res["seq_len"],
        name=t.get("name"),
        max=t.get("max_s"),
        total=t.get("total_s"),
        query_cover=t.get("qc_s"),
        evalue=t.get("e_s"),
        per_ident=t.get("pi_s"),
        acc_len=t.get("acc_len"),
        accession=t.get("accession"),
        taxid=t.get("taxid"),
        lineage=res.get("lineage"),
        lineage_full=res.get("lineage_full"),
        flag=level or "",
        comment=comment,
        db=res.get("db"),
        params=res.get("params"),
        blast_version=res.get("blast_version"),
        elapsed_s=round(res.get("elapsed") or 0, 1),
        verified="",
        source=res.get("source"),
        search_backend=res.get("backend") or "remote",
        local_db=res.get("local_db") or "",
        top15=[
            {k: h[k] for k in ("name", "accession", "max_s", "total_s", "qc_s", "e_s", "pi_s", "acc_len")}
            for h in res["hits"][:15]
        ],
    )
    if merge is not None:
        rec.update(
            merge_mode=merge.mode,
            merge_backend=merge.backend,
            direction=merge.direction,
            overlap_len=merge.overlap_len,
            overlap_ident=merge.overlap_ident,
            n_conflicts=merge.n_conflicts,
            n_ambiguities=merge.n_ambiguities,
            trimmed_len=merge.trimmed_len,
        )
    return rec


def write_logs(outdir: Path, records: list[dict]) -> None:
    (outdir / "log.json").write_text(json.dumps(records, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    with open(outdir / "log.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=LOG_FIELDS, extrasaction="ignore")
        w.writeheader()
        for rec in records:
            w.writerow(rec)


def _close_number(cell, expected, *, evalue=False) -> bool:
    if cell is None and expected is None:
        return True
    if evalue and (expected == 0 or expected == 0.0):
        try:
            return float(cell) == 0.0
        except (TypeError, ValueError):
            return False
    try:
        a, b = float(cell), float(expected)
    except (TypeError, ValueError):
        return cell == expected
    if evalue:
        if a == 0 or b == 0:
            return a == b
        return abs(a - b) / max(abs(a), abs(b)) <= 1e-6
    return abs(a - b) <= 1e-8


def _rerun_local_entries(entries, ws, template, field_cols, log):
    """로컬 행의 서열을 시트에서 읽어 BLAST+를 다시 실행한다. 값은 행 번호 → 결과 또는 오류 문자열."""
    if not entries:
        return {}
    from blastid.localblast import run_local

    grouped: dict[str, list] = {}
    out = {}
    seq_col = field_cols.get(template.sequence_field)
    for entry in entries:
        row = entry["row"]
        if not seq_col:
            out[row] = "서열 열이 없어 로컬 재검증을 할 수 없습니다"
            continue
        seq = clean_seq(ws.cell(row, seq_col).value, 1)
        if not seq:
            out[row] = "시트에서 서열을 읽지 못했습니다"
            continue
        db = entry.get("local_db") or "ITS_RefSeq_Fungi"
        grouped.setdefault(db, []).append((entry, seq))
    for db, group in grouped.items():
        queries = [
            dict(key=entry["row"], label=f"row{entry['row']}", seq=seq)
            for entry, seq in group
        ]
        try:
            import tempfile

            found = run_local(queries, db, Path(tempfile.mkdtemp(prefix="blastid-verify-")), log)
        except Exception as exc:
            for entry, _seq in group:
                out[entry["row"]] = str(exc)
            continue
        for entry, seq in group:
            result = found.get(entry["row"])
            if not result or not result["hits"]:
                out[entry["row"]] = None
                continue
            top = result["hits"][0]
            out[entry["row"]] = dict(
                qlen=result["qlen"],
                name=top["name"],
                max=top["max"],
                total=top["total"],
                qc=top["qc"],
                evalue=top["evalue"],
                pi=top["pi"],
                acc_len=top["acc_len"],
                accession=top["accession"],
                taxid=top.get("taxid"),
            )
    return out


def _is_local_entry(entry: dict) -> bool:
    return entry.get("backend") == "local" or str(entry.get("rid") or "").startswith("local")


def verify(ncbi, ws, template: Template, field_cols: dict[str, int], entries: list[dict], log=print, pause: float = 1.0):
    """저장된 시트를 RID XML·Taxonomy 재조회로 셀마다 비교. 로컬 행은 BLAST+를 다시 돌린다."""
    remote_entries = [e for e in entries if not _is_local_entry(e)]
    local_entries = [e for e in entries if _is_local_entry(e)]
    xml_cache = {}
    for rid in sorted({e["rid"] for e in remote_entries}):
        xml_cache[rid] = parse_xml_report(ncbi.get_xml(rid))
        if pause:
            time.sleep(pause)
    local_cache = _rerun_local_entries(local_entries, ws, template, field_cols, log)
    tax = ncbi.taxonomy(e["taxid"] for e in remote_entries if e.get("taxid"))
    local_taxids = [item["taxid"] for item in local_cache.values() if isinstance(item, dict) and item.get("taxid")]
    if local_taxids:
        tax.update(ncbi.taxonomy(local_taxids))
    checks = [
        ("scientific_name", "name", "str"),
        ("max_score", "max", "num"),
        ("total_score", "total", "num"),
        ("query_cover", "qc", "num"),
        ("evalue", "evalue", "evalue"),
        ("per_ident", "pi", "num"),
        ("acc_len", "acc_len", "num"),
        ("accession", "accession", "str"),
        ("lineage", "lineage", "str"),
    ]
    problems = {}
    for e in entries:
        r = e["row"]
        bad = []
        if _is_local_entry(e):
            x = local_cache.get(r)
            if isinstance(x, str):
                problems[r] = [x]
                log(f"  검증 {r}행: 불일치 {x}")
                continue
        else:
            x = xml_cache[e["rid"]][e["query_index"] - 1] if e["query_index"] - 1 < len(xml_cache[e["rid"]]) else None
        name_col = field_cols.get("scientific_name")
        cell_name = ws.cell(r, name_col).value if name_col else None
        if x is None:
            if name_col and cell_name != "No significant similarity found":
                bad.append("hit 없음인데 Scientific Name 이 다름")
        else:
            if x["qlen"] != e["seq_len"]:
                bad.append(f"query 길이 {x['qlen']} ≠ 서열 {e['seq_len']}")
            shown_name = x.get("name") if _is_local_entry(e) and x.get("name") else e.get("name_json")
            expected = {
                "name": shown_name,
                "max": x["max"],
                "total": x["total"],
                "qc": x["qc"],
                "evalue": x["evalue"],
                "pi": x["pi"],
                "acc_len": x["acc_len"],
                "accession": x["accession"],
            }
            taxid = x.get("taxid") if _is_local_entry(e) else e.get("taxid")
            tinfo = tax.get(str(taxid)) if taxid else None
            if "lineage" in field_cols:
                cell_lineage = ws.cell(r, field_cols["lineage"]).value
                if not taxid:
                    if cell_lineage not in (None, ""):
                        bad.append("taxid가 없어 lineage를 다시 확인하지 못했습니다")
                elif not tinfo:
                    bad.append("taxonomy 조회 실패")
                else:
                    expected["lineage"] = fungi_lineage(tinfo["lineage_full"])
                    if tinfo["name"] != shown_name:
                        log(
                            f"  (참고) {r}행: BLAST 표시명 '{shown_name}' "
                            f"vs Taxonomy 현재명 '{tinfo['name']}'"
                        )
            for fld, key, kind in checks:
                col = field_cols.get(fld)
                if not col or key not in expected:
                    continue
                got = ws.cell(r, col).value
                exp = expected[key]
                if kind == "str":
                    if (got or "") != (exp or ""):
                        bad.append(f"{fld}: 엑셀 {got!r} ≠ 재계산 {exp!r}")
                elif not _close_number(got, exp, evalue=(kind == "evalue")):
                    bad.append(f"{fld}: 엑셀 {got!r} ≠ 재계산 {exp!r}")
        if bad:
            problems[r] = bad
        log(f"  검증 {r}행: {'OK' if not bad else '불일치 ' + '; '.join(bad)}")
    return problems


def _reuse(args):
    raw = getattr(args, "reuse_rids", None)
    if not raw:
        return None
    return [x.strip() for x in str(raw).split(",") if x.strip()]


def _bind_columns(ws, template: Template, args, log):
    cols = map_fields(ws, template)
    if getattr(args, "seq_header", None):
        found = None
        from blastid.workbook import header_map

        hm = header_map(ws, template.header_row)
        key = norm_header(args.seq_header)
        if key not in hm:
            raise SystemExit(f"--seq-header '{args.seq_header}' 열을 찾지 못했습니다.")
        cols[template.sequence_field] = hm[key]
    if getattr(args, "id_header", None):
        from blastid.workbook import header_map

        hm = header_map(ws, template.header_row)
        key = norm_header(args.id_header)
        if key not in hm:
            raise SystemExit(f"--id-header '{args.id_header}' 열을 찾지 못했습니다.")
        cols[template.id_field] = hm[key]
    if template.sequence_field not in cols:
        guessed = guess_sequence_column(ws, template.header_row, getattr(args, "min_seq_len", 50))
        if guessed:
            cols[template.sequence_field] = guessed
            log(f"서열 열을 헤더가 아니라 셀 내용으로 찾았습니다 (열 {guessed}).")
        else:
            raise SystemExit(
                "서열 열을 찾지 못했습니다. 템플릿 header/aliases 또는 --seq-header 를 지정하세요. "
                f"첫 행: {[ws.cell(template.header_row, c).value for c in range(1, (ws.max_column or 0) + 1)]}"
            )
    if template.id_field not in cols:
        guessed = guess_id_column(ws, template.header_row, cols.get(template.sequence_field))
        if guessed:
            cols[template.id_field] = guessed
            log(f"ID 열을 헤더로 추정했습니다 (열 {guessed}).")
    return cols


def _excel_queries(ws, template, cols, args, log):
    id_col = cols.get(template.id_field)
    seq_col = cols[template.sequence_field]
    name_col = cols.get(template.empty_field)
    bold_col = cols.get(template.bold_field, id_col)
    header = template.header_row
    if getattr(args, "rows", None):
        rows = parse_rows(args.rows)
    else:
        rows = list(range(header + 1, (ws.max_row or header) + 1))
    queries = []
    n_short = n_empty = n_bold = 0
    min_len = getattr(args, "min_seq_len", 50)
    for r in rows:
        id_val = ws.cell(r, id_col).value if id_col else None
        raw_seq = ws.cell(r, seq_col).value
        if id_val in (None, "") and raw_seq in (None, ""):
            continue
        if getattr(args, "bold_only", False) and bold_col:
            font = ws.cell(r, bold_col).font
            if not (font and (font.bold or font.b)):
                n_bold += 1
                continue
        if getattr(args, "only_empty", False) and name_col:
            cur = ws.cell(r, name_col).value
            if cur not in (None, ""):
                n_empty += 1
                continue
        seq = clean_seq(raw_seq, min_len)
        if not seq:
            n_short += 1
            log(f"건너뜀 {r}행 {id_val}: 서열이 아니거나 {min_len}bp 미만")
            continue
        sid = str(id_val).strip() if id_val not in (None, "") else f"row{r}"
        queries.append(
            dict(key=r, row=r, label=fasta_label(sid, f"row{r}"), sample_id=sid, seq=seq, merge=None)
        )
    unique_labels(queries)
    log(
        f"대상 {len(queries)}행"
        + (f" ({', '.join(str(q['sample_id']) for q in queries)})" if queries else "")
        + f"  / 건너뜀 서열 {n_short}, 이미 채움 {n_empty}, 굵게 아님 {n_bold}"
    )
    return queries


def _apply(wb, ws, template, cols, queries, results, args, outdir, log, *, write_inputs: bool):
    shots = Path(outdir) / "shots"
    shots.mkdir(exist_ok=True)
    records, entries, flags = [], [], {}
    comment_col = cols.get(template.comment_field)
    for q in queries:
        rr = results[q["key"]]
        blast_level, blast_comment = assess(
            rr["hits"], args.min_ident, args.min_cover, args.alt_margin
        )
        level, comment = combine_review(blast_level, blast_comment, q.get("merge"))
        if rr.get("backend") == "local":
            note = local_note(rr.get("db") or "local")
            comment = f"{comment} {note}".strip() if comment else note
        elif comment:
            comment = f"{comment} (RID {rr['rid']})"
        row = q["row"]
        if level:
            flags[row] = level
        elif comment_col and not write_inputs:
            cell = ws.cell(row, comment_col)
            if not comment and should_clear_comment(cell, getattr(args, "clear_old_comments", False)):
                clear_review(ws, row, template, cols)
            elif not comment:
                old = level_from_comment(cell.value)
                if old:
                    flags[row] = old
                    log(f"{row}행: 새 판정은 깨끗하지만 기존 코멘트를 남겨 두었습니다.")
        values = build_field_values(
            no=q.get("no"),
            sample_id=q["sample_id"],
            sequence=q["seq"],
            blast=rr,
            level=level,
            comment=comment,
            merge=q.get("merge"),
        )
        write_fields(ws, row, template, cols, values, fill_input=write_inputs)
        if level:
            safe = re.sub(r"[^\w.\-+]+", "_", q["sample_id"])[:40]
            png = shots / f"row{row}_{safe}.png"
            try:
                from blastid.render import render_table

                render_table(
                    png,
                    f"{row}행 {q['sample_id']}  |  {rr.get('source') or rr['rid']}  |  query {rr['qlen']} bp",
                    rr["hits"],
                )
                place_shot(wb, ws, row, png, f"{row}행 {q['sample_id']} — {comment}", template, cols)
            except Exception as e:
                log(f"  스크린샷 실패 {row}행: {e}")
        rec = make_record(row, q["sample_id"], rr, level, comment, q.get("merge"))
        records.append(rec)
        entries.append(
            dict(
                row=row,
                rid=rr["rid"],
                query_index=rr["query_index"],
                seq_len=len(q["seq"]),
                taxid=rr["hits"][0]["taxid"] if rr["hits"] else None,
                name_json=rr["hits"][0]["name"] if rr["hits"] else None,
                backend=rr.get("backend") or "remote",
                local_db=rr.get("local_db") or "",
            )
        )
        mark = f"  ⚠ {level}" if level else ""
        log(
            f"{row}행 {q['sample_id']}: {rec['name']} | {rec['max']}/{rec['total']} | "
            f"{rec['query_cover']} | {rec['evalue']} | {rec['per_ident']} | "
            f"{rec['acc_len']} | {rec['accession']}{mark}"
        )
    if not write_inputs and comment_col:
        done = {q["row"] for q in queries}
        for r in range(template.header_row + 1, (ws.max_row or 1) + 1):
            if r in flags or r in done:
                continue
            old = level_from_comment(ws.cell(r, comment_col).value)
            if old:
                flags[r] = old
    local_rows = {q["row"] for q in queries if results[q["key"]].get("backend") == "local"}
    style_sheet(ws, template, cols, flags, local_rows)
    return records, entries


def _finish(ncbi, wb, out, sheet_name, template, cols, entries, records, args, outdir, log, t0, n):
    wb.save(out)
    log(f"엑셀 저장: {out}  (지금까지 {time.time() - t0:.0f}초, 서열당 {(time.time() - t0) / max(n, 1):.1f}초)")
    if not getattr(args, "no_verify", False) and entries:
        log("독립 검증: 저장한 파일을 다시 열고 RID 결과(XML)와 Taxonomy를 새로 받아 셀마다 비교")
        import openpyxl

        ws2 = pick_sheet(openpyxl.load_workbook(out), sheet_name, log)
        pause = 0.0 if getattr(ncbi, "fast", False) else 1.0
        probs = verify(ncbi, ws2, template, cols, entries, log, pause=pause)
        for rec in records:
            rec["verified"] = rec["row"] not in probs
        log(
            f"검증 결과: {len(entries) - len(probs)}/{len(entries)}행 일치"
            + (f", 불일치 {sorted(probs)}" if probs else "")
        )
    write_logs(Path(outdir), records)
    log(f"로그: {outdir / 'log.csv'}, {outdir / 'log.json'}")


def _job_email(args):
    """None 이면 환경변수. 문자열(빈 문자열 포함)이면 그 값만 사용."""
    if "email" in getattr(args, "__dict__", {}):
        return args.email
    return None


def blast_queries(queries, args, outdir, log, ncbi=None):
    if not queries:
        log("BLAST할 서열이 없습니다.")
        return None, None, None
    reuse = _reuse(args)
    batch = max(1, int(getattr(args, "batch_size", 10) or 10))
    usage = UsageLog()
    backend = getattr(args, "backend", "auto") or "auto"
    limit = int(getattr(args, "remote_limit", 100) or 100)
    n_batches = 0 if reuse else math.ceil(len(queries) / batch)
    n_searches = remote_batch_budget(usage.count(), n_batches, limit, "remote" if reuse else backend)
    plan = plan_submission(
        n_searches,
        bool(getattr(args, "off_hours", False)),
        dt.datetime.now(ET_ZONE),
        usage,
    )
    for warning in plan.warnings:
        log("NCBI: " + warning)
    if getattr(args, "dry_run", False):
        for q in queries:
            log(f"[dry-run] {q['sample_id']}  {len(q['seq'])} bp")
        log("dry-run: NCBI에 제출하지 않았습니다.")
        return None, None, None
    if plan.wait_until is not None:
        sleep_until(plan.wait_until, log)
    ncbi = ncbi or NCBI(log, email=_job_email(args))
    if not ncbi.email:
        log("이메일을 보내지 않습니다. NCBI_EMAIL 또는 작업창의 이메일이 비어 있습니다.")
    log(DISCLAIMER_KO)
    try:
        res = run_blast(
            ncbi,
            queries,
            batch,
            outdir,
            log,
            reuse,
            usage=usage,
            backend=getattr(args, "backend", "auto") or "auto",
            remote_limit=int(getattr(args, "remote_limit", 100) or 100),
            local_db=getattr(args, "local_db", None) or "ITS_RefSeq_Fungi",
            blastn=getattr(args, "blastn", None),
            max_wait=float(getattr(args, "max_wait", 60) or 60) * 60.0,
        )
    except LocalBlastUnavailable as exc:
        partial = exc.partial or {}
        if partial:
            check_params(partial, log)
            add_taxonomy(ncbi, partial, log)
        log(str(exc))
        return ncbi, partial, str(exc)
    check_params(res, log)
    add_taxonomy(ncbi, res, log)
    return ncbi, res, None


def run_excel(args, template: Template, outdir: Path, log):
    import openpyxl

    src = Path(args.inputs[0]).resolve()
    out = Path(args.output).resolve() if args.output else None
    if out is None:
        out = src if getattr(args, "in_place", False) else Path(outdir) / f"{src.stem}_blast.xlsx"
    if out == src and not getattr(args, "in_place", False) and src.suffix.lower() != ".csv":
        raise SystemExit("출력이 입력 파일과 같습니다. 덮어쓰려면 --in-place 를 쓰세요.")
    if src.suffix.lower() == ".csv":
        wb = workbook_from_csv(src)
        if out is None or out == src:
            out = Path(outdir) / f"{src.stem}_blast.xlsx"
            log(f"CSV는 서식·그림을 담을 수 없어 xlsx로 저장합니다: {out}")
    else:
        wb = openpyxl.load_workbook(src)
    sheet_name = args.sheet or template.sheet
    ws = pick_sheet(wb, sheet_name, log)
    cols = _bind_columns(ws, template, args, log)
    cols = ensure_fields(ws, template, cols, log)
    queries = _excel_queries(ws, template, cols, args, log)
    t0 = time.time()
    ncbi, res, stopped = blast_queries(queries, args, outdir, log)
    if res is None:
        return
    done = [q for q in queries if q["key"] in res]
    if done:
        records, entries = _apply(wb, ws, template, cols, done, res, args, outdir, log, write_inputs=False)
        _finish(ncbi, wb, out, ws.title, template, cols, entries, records, args, outdir, log, t0, len(done))
    if stopped:
        raise SystemExit(stopped)


def run_fresh(queries, args, template: Template, outdir: Path, log):
    """FASTA 또는 합친 ab1 서열로 템플릿 워크북을 새로 만든다."""
    if not queries:
        log("BLAST할 서열이 없습니다.")
        return
    t0 = time.time()
    ncbi, res, stopped = blast_queries(queries, args, outdir, log)
    if res is None:
        return
    queries = [q for q in queries if q["key"] in res]
    if not queries:
        if stopped:
            raise SystemExit(stopped)
        return
    wb = new_workbook(template)
    sheet_name = args.sheet or template.sheet
    if sheet_name not in wb.sheetnames and wb.sheetnames:
        # 템플릿 파일의 시트 이름이 YAML과 다를 수 있다. 시트가 하나면 그 시트를 쓴다.
        if len(wb.sheetnames) == 1:
            wb.active.title = sheet_name if sheet_name else wb.sheetnames[0]
    ws = pick_sheet(wb, sheet_name, log)
    cols = map_fields(ws, template)
    if template.sequence_field not in cols or template.empty_field not in cols:
        # 템플릿 xlsx 가 비어 있거나 헤더가 다르면 YAML 순서로 다시 만든다.
        wb = __import__("blastid.template", fromlist=["build_workbook"]).build_workbook(template)
        ws = wb[template.sheet]
        cols = map_fields(ws, template)
    cols = ensure_fields(ws, template, cols, log)
    records, entries = _apply(wb, ws, template, cols, queries, res, args, outdir, log, write_inputs=True)
    out = Path(args.output) if args.output else Path(outdir) / "blast_results.xlsx"
    _finish(ncbi, wb, out, ws.title, template, cols, entries, records, args, outdir, log, t0, len(queries) or 1)
    if stopped:
        raise SystemExit(stopped)


def load_fasta_queries(paths, min_len, log) -> list[dict]:
    from Bio import SeqIO

    queries = []
    n = 0
    for p in paths:
        for rec in SeqIO.parse(str(p), "fasta"):
            n += 1
            seq = clean_seq(str(rec.seq), min_len)
            sid = rec.id or f"seq{n}"
            if not seq:
                log(f"건너뜀 {sid}: 서열이 아니거나 {min_len}bp 미만")
                continue
            queries.append(
                dict(
                    key=len(queries),
                    label=fasta_label(sid, f"seq{n}"),
                    sample_id=sid,
                    seq=seq,
                    merge=None,
                )
            )
    for i, q in enumerate(queries):
        q["key"] = i
        q["row"] = i + 2
        q["no"] = i + 1
    unique_labels(queries)
    log(f"FASTA {len(queries)}개")
    return queries


def load_ab1_queries(args, log) -> list[dict]:
    regexes = getattr(args, "pair_regex", None) or None
    results = assemble_files(
        args.inputs,
        pairs_csv=getattr(args, "pairs", None),
        pair_regexes=regexes,
        no_merge=getattr(args, "no_merge", False),
        trimmer=getattr(args, "trimmer", "mott"),
        trim_cutoff=getattr(args, "trim_cutoff", 0.05),
        min_len=getattr(args, "min_len", 150),
        min_q=getattr(args, "min_q", 25),
        min_overlap=getattr(args, "min_overlap", 50),
        min_overlap_ident=getattr(args, "min_overlap_ident", 90.0),
        conflict_q=getattr(args, "conflict_q", 5),
        indel_q=getattr(args, "indel_q", 15),
        merger=getattr(args, "merger", "biopython"),
        tracy_trim=getattr(args, "tracy_trim", 4),
        log=log,
    )
    queries = []
    for i, res in enumerate(results):
        if not res.seq:
            continue
        queries.append(
            dict(
                key=i,
                row=i + 2,
                no=i + 1,
                label=fasta_label(res.sample_id, f"seq{i+1}"),
                sample_id=res.sample_id,
                seq=res.seq,
                merge=res,
            )
        )
    unique_labels(queries)
    # key 를 다시 0.. 로 맞춘다. unique_labels 는 label만 바꾼다.
    for i, q in enumerate(queries):
        q["key"] = i
    return results, queries


def detect_kind(paths: list[str]) -> str:
    kinds = []
    for raw in paths:
        path = Path(raw)
        if path.is_dir():
            kinds.append("ab1")
            continue
        suf = path.suffix.lower()
        if suf in {".ab1", ".abi"}:
            kinds.append("ab1")
        elif suf in {".fasta", ".fa", ".fna", ".ffn", ".fas"}:
            kinds.append("fasta")
        elif suf in {".xlsx", ".xlsm", ".csv"}:
            kinds.append("excel")
        elif not path.exists():
            raise SystemExit(f"입력이 없습니다: {path}")
        else:
            head = path.read_bytes()[:8]
            if head.startswith(b"ABIF"):
                kinds.append("ab1")
            elif head.startswith(b">"):
                kinds.append("fasta")
            elif head.startswith(b"PK"):
                kinds.append("excel")
            else:
                raise SystemExit(f"입력 종류를 알 수 없습니다: {path}")
    if not kinds:
        raise SystemExit("입력이 없습니다.")
    if len(set(kinds)) != 1:
        raise SystemExit("엑셀, FASTA, ab1 을 한 번에 섞지 마세요.")
    if kinds[0] == "excel" and len(paths) != 1:
        raise SystemExit("엑셀·CSV는 파일을 하나만 지정하세요.")
    return kinds[0]


def read_text_auto(path: Path) -> str:
    """UTF-8(BOM 허용)을 먼저 보고, 아니면 CP949로 읽는다."""
    raw = Path(path).read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp949")


def workbook_from_csv(path: Path):
    """CSV를 첫 시트가 표인 워크북으로 만든다. 구분자 , ; 탭을 가늠한다."""
    import openpyxl

    text = read_text_auto(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value if value != "" else None)
    return wb


def execute(args) -> None:
    outdir = default_outdir(getattr(args, "out_dir", None))
    log = make_logger(outdir)
    kind = args.kind
    if kind == "run":
        kind = detect_kind(args.inputs)
        if kind == "ab1":
            args.blast = True
    template = load_template(getattr(args, "template", None))
    if getattr(args, "sheet", None):
        template.sheet = args.sheet
    log(f"blastid {kind}  템플릿 {template.source}")
    if kind == "excel":
        run_excel(args, template, outdir, log)
        return
    if kind == "fasta":
        queries = load_fasta_queries(args.inputs, getattr(args, "min_seq_len", 50), log)
        run_fresh(queries, args, template, outdir, log)
        return
    if kind == "ab1":
        merged, queries = load_ab1_queries(args, log)
        write_qc(outdir, merged)
        log(f"FASTA: {outdir / 'merged.fasta'}  QC: {outdir / 'ab1_qc.csv'}, {outdir / 'merge_qc.csv'}")
        if not getattr(args, "blast", False):
            log("BLAST는 하지 않았습니다. 바로 동정하려면 --blast 또는 `blastid run` 을 쓰세요.")
            return
        run_fresh(queries, args, template, outdir, log)
        return
    raise SystemExit(f"알 수 없는 명령: {kind}")
