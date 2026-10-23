from pathlib import Path

import openpyxl

from blastid.cli import main


def test_help_and_version(capsys):
    try:
        main(["--help"])
    except SystemExit as e:
        assert e.code in (0, None)
    out = capsys.readouterr().out
    assert "run" in out and "init-template" in out


def test_init_template_command(tmp_path, capsys):
    main(["init-template", "-o", str(tmp_path / "out")])
    assert (tmp_path / "out" / "template.yaml").exists()
    assert (tmp_path / "out" / "template.xlsx").exists()
    text = capsys.readouterr().out
    assert "template.yaml" in text


def test_path_without_subcommand_is_run(tmp_path, capsys):
    seq = "ACGT" * 20
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ID(Voucher)"
    ws["B1"] = "Seq"
    ws["A2"] = "S1"
    ws["B2"] = seq
    src = tmp_path / "in.xlsx"
    wb.save(src)
    main([str(src), "--dry-run", "--out-dir", str(tmp_path / "runs")])
    log = capsys.readouterr().out
    assert "dry-run" in log
    assert "S1" in log
