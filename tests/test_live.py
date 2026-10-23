"""선택적 실 NCBI 확인. 기본 pytest 에서는 빠진다.

    BLASTID_LIVE_XLSX=/path/to/sample.xlsx pytest -m live
"""
import os

import openpyxl
import pytest

from blastid.cli import main

# 손으로 맞춘 1위. DB가 바뀌면 학명·accession·퍼센트가 조금 달라질 수 있다.
EXPECTED = {
    "OhF260617-02": {"name": "Pyxine copelandii", "acc": "EU670225.1", "ident": 100.00},
    "OhF260623-14": {"name": "Lecidella biformis", "acc": "PZ322662.1", "ident": 87.21, "cover": 96},
    "OhF260625-05": {"name": "Trimmatothelopsis sp.", "ident": 94.86},
    "OhF260630-26": {"name": "Collema furfuraceum", "ident": 88.82},
    "OhF260721-17": {"name": "Protoblastenia cyclospora", "acc": "AY425668.1", "ident": 79.96},
}


@pytest.mark.live
def test_live_sample_top_hits(tmp_path):
    src = os.environ.get("BLASTID_LIVE_XLSX")
    if not src:
        pytest.skip("BLASTID_LIVE_XLSX 가 없습니다")
    out = tmp_path / "live.xlsx"
    main(
        [
            "run",
            src,
            "-o",
            str(out),
            "--out-dir",
            str(tmp_path / "runs"),
            "--rows",
            "2-6",
        ]
    )
    ws = openpyxl.load_workbook(out)["Sheet1"]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    idc, namec, identc = headers["ID(Voucher)"], headers["Scientific Name"], headers["Per. Ident"]
    accc, covc = headers["Accession"], headers["Query Cover"]
    found = {}
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(r, idc).value
        if sid in EXPECTED:
            found[sid] = {
                "name": ws.cell(r, namec).value,
                "acc": ws.cell(r, accc).value,
                "ident": round(float(ws.cell(r, identc).value) * 100, 2),
                "cover": round(float(ws.cell(r, covc).value) * 100),
            }
    assert set(found) == set(EXPECTED)
    for sid, exp in EXPECTED.items():
        got = found[sid]
        assert got["name"] == exp["name"], sid
        if "acc" in exp:
            assert got["acc"] == exp["acc"], sid
        assert got["ident"] == pytest.approx(exp["ident"], abs=0.05), sid
        if "cover" in exp:
            assert got["cover"] == exp["cover"], sid
