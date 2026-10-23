"""가짜 NCBI 응답으로 엑셀 기록·검증·서식을 확인한다."""
import io
from pathlib import Path
from urllib.request import Request

import openpyxl
import pytest
from openpyxl.styles import Font

from blastid.fields import build_field_values
from blastid.ncbi import NCBI, blast_identity_params
from blastid.pipeline import execute, run_blast
from blastid.scoring import web_values
from blastid.template import load_template
from blastid.workbook import ensure_fields, map_fields, style_sheet, write_fields
from tests.blastlib import blast_json, blast_xml, json_hit


def _spec(name, acc, ident, alen, qlen, bit=500.2, evalue=0.0, taxid=42, length=800, qto=None):
    return {
        "name": name,
        "taxid": taxid,
        "acc": acc,
        "title": f"{name} ITS",
        "length": length,
        "hsps": [
            {
                "bit": bit,
                "ident": ident,
                "alen": alen,
                "qfrom": 1,
                "qto": qlen if qto is None else qto,
                "evalue": evalue,
            }
        ],
    }


class FakeNCBI:
    def __init__(self, queries):
        self.email = ""
        self.api_key = ""
        self.fast = True
        self._last_poll = {}
        self.submitted = []
        self.queries = queries
        self.json = blast_json(queries)
        self.xml = blast_xml(queries)
        self.tax = {
            "42": {
                "taxid": "42",
                "name": "Kept Name",
                "lineage_full": "cellular organisms; Eukaryota; Fungi; Dikarya",
            },
            "7": {
                "taxid": "7",
                "name": "Other",
                "lineage_full": "cellular organisms; Fungi; Dikarya; Ascomycota",
            },
        }

    def submit(self, fasta):
        self.submitted.append(fasta)
        return "FAKERID", 0

    def wait_all(self, jobs):
        import time

        return {rid: time.time() for rid in jobs}

    def get_json(self, rid):
        assert rid == "FAKERID"
        return self.json

    def get_xml(self, rid):
        return self.xml

    def taxonomy(self, taxids):
        return self.tax


def test_excel_maps_by_header_not_column_letter(tmp_path, monkeypatch):
    seq_ok = "ACGT" * 30  # 120
    seq_low = "ACGT" * 25  # 100
    queries = [
        ("Good_1", seq_ok, [_spec("Pyxine copelandii", "EU670225.1", 120, 120, 120, bit=900.4)]),
        ("Low_2", seq_low, [_spec("Lecidella biformis", "PZ322662.1", 80, 100, 100, bit=300.9, evalue=1.23e-50)]),
    ]
    fake = FakeNCBI(queries)
    monkeypatch.setattr("blastid.pipeline.NCBI", lambda log=None: fake)

    # 열 순서를 일부러 섞는다. K열이 Scientific Name 이 아니어야 한다.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["Seq", "메모", "ID(Voucher)", "Per. Ident", "Scientific Name", "Query Cover", "Accession", "E value",
               "Max Score", "Total Score", "Acc. Len", "Lineage", "검토 코멘트", "스크린샷 link"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    ws.cell(2, 3, "Good_1")
    ws.cell(2, 1, seq_ok)
    ws.cell(2, 2, "남길 메모")
    ws.cell(3, 3, "Low_2")
    ws.cell(3, 1, seq_low)
    ws.cell(4, 3, "SkipMe")
    ws.cell(4, 5, "이미 있음")
    ws.cell(4, 1, seq_ok)
    src = tmp_path / "in.xlsx"
    wb.save(src)
    out = tmp_path / "out.xlsx"
    execute(
        _args(
            kind="excel",
            inputs=[str(src)],
            output=str(out),
            out_dir=str(tmp_path / "run"),
            only_empty=True,
        )
    )
    assert fake.submitted and "SkipMe" not in fake.submitted[0]
    assert "Good_1" in fake.submitted[0]
    saved = openpyxl.load_workbook(out)
    ws = saved["Sheet1"]
    # Scientific Name 은 5열, Seq 는 1열. 글자 위치 그대로.
    assert ws.cell(1, 5).value == "Scientific Name"
    assert ws.cell(2, 5).value == "Pyxine copelandii"
    assert ws.cell(2, 7).value == "EU670225.1"
    assert ws.cell(2, 1).value == seq_ok
    assert ws.cell(2, 2).value == "남길 메모"
    assert ws.cell(2, 9).value == 900  # 900.4 내림
    assert ws.cell(2, 6).value == 1  # 100% cover, 엑셀은 1
    assert ws.cell(2, 6).number_format == "0%"
    assert ws.cell(2, 4).number_format == "0.00%"
    assert ws.cell(2, 4).value == 1
    assert ws.cell(2, 8).value == 0
    assert ws.cell(2, 12).value.startswith("Fungi;")
    assert ws.cell(3, 5).value == "Lecidella biformis"
    assert ws.cell(3, 7).value == "PZ322662.1"
    assert ws.cell(3, 4).value == 0.8
    assert str(ws.cell(3, 13).value).startswith("[이상]")
    assert ws.cell(3, 14).hyperlink is not None
    assert ws.cell(4, 5).value == "이미 있음"  # only-empty
    assert saved["스크린샷"]._images
    assert ws.freeze_panes == "D2"
    assert ws.row_dimensions[2].height == 30
    assert ws.auto_filter.ref.startswith("A1:")
    rgb = ws.cell(1, 5).fill.fgColor.rgb
    assert rgb.endswith("1F4E78")
    # 낮은 일치도 행은 빨강
    assert ws.cell(3, 5).fill.fgColor.rgb.endswith("F8CBAD")
    # 검증이 통과해 log 에 verified True
    import json

    logs = list((tmp_path / "run").glob("*/log.json"))
    recs = json.loads(logs[0].read_text(encoding="utf-8"))
    assert recs and all(r["verified"] is True for r in recs)
    assert {r["id"] for r in recs} == {"Good_1", "Low_2"}


def test_fresh_workbook_from_fasta(tmp_path, monkeypatch):
    seq = "ACGT" * 40
    queries = [("SampleA", seq, [_spec("Pyxine copelandii", "EU670225.1", len(seq), len(seq), len(seq))])]
    monkeypatch.setattr("blastid.pipeline.NCBI", lambda log=None: FakeNCBI(queries))
    fasta = tmp_path / "s.fasta"
    fasta.write_text(f">SampleA\n{seq}\n", encoding="utf-8")
    out = tmp_path / "res.xlsx"
    execute(_args(kind="fasta", inputs=[str(fasta)], output=str(out), out_dir=str(tmp_path / "run")))
    ws = openpyxl.load_workbook(out)["Sheet1"]
    headers = [ws.cell(1, c).value for c in range(1, 22)]
    assert headers[2] == "ID(Voucher)"
    assert headers[9] == "Seq"
    assert headers[10] == "Scientific Name"
    assert ws.cell(2, 3).value == "SampleA"
    assert ws.cell(2, 11).value == "Pyxine copelandii"
    assert ws.freeze_panes == "D2"


def test_dry_run_does_not_call_ncbi(tmp_path, monkeypatch):
    def boom(*_a, **_k):
        raise AssertionError("NCBI가 호출되면 안 됩니다")

    monkeypatch.setattr("blastid.pipeline.NCBI", boom)
    seq = "ACGT" * 20
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "ID(Voucher)")
    ws.cell(1, 2, "Seq")
    ws.cell(2, 1, "S")
    ws.cell(2, 2, seq)
    src = tmp_path / "t.xlsx"
    wb.save(src)
    execute(_args(kind="excel", inputs=[str(src)], out_dir=str(tmp_path / "run"), dry_run=True))


def test_bold_only(tmp_path, monkeypatch):
    seq = "ACGT" * 20
    captured = {}

    class One(FakeNCBI):
        def __init__(self):
            super().__init__([("Bold", seq, [_spec("Pyxine copelandii", "EU670225.1", len(seq), len(seq), len(seq))])])

        def submit(self, fasta):
            captured["fasta"] = fasta
            return super().submit(fasta)

    monkeypatch.setattr("blastid.pipeline.NCBI", lambda log=None: One())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ID(Voucher)"
    ws["B1"] = "Seq"
    ws["C1"] = "Scientific Name"
    ws["A2"] = "Bold"
    ws["A2"].font = Font(bold=True)
    ws["B2"] = seq
    ws["A3"] = "Plain"
    ws["B3"] = seq
    src = tmp_path / "b.xlsx"
    wb.save(src)
    out = tmp_path / "o.xlsx"
    execute(
        _args(
            kind="excel",
            inputs=[str(src)],
            output=str(out),
            out_dir=str(tmp_path / "run"),
            bold_only=True,
            no_verify=True,
        )
    )
    assert "Bold" in captured["fasta"] and "Plain" not in captured["fasta"]
    ws = openpyxl.load_workbook(out).active
    assert ws["C2"].value == "Pyxine copelandii"
    assert ws["C3"].value is None


def test_blast_request_omits_email_unless_env(monkeypatch):
    monkeypatch.delenv("NCBI_EMAIL", raising=False)
    monkeypatch.delenv("NCBI_API_KEY", raising=False)
    assert blast_identity_params() == [("tool", "blastid")]
    seen = {}

    class Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_open(req: Request, timeout=0):
        seen["url"] = req.full_url
        seen["data"] = req.data.decode() if req.data else ""
        return Resp(b"RID = ABC\nRTOE = 0\n")

    monkeypatch.setattr("blastid.ncbi.urllib.request.urlopen", fake_open)
    ncbi = NCBI(log=lambda *_: None)
    rid, _rtoe = ncbi.submit(">q\n" + "ACGT" * 20 + "\n")
    assert rid == "ABC"
    body = seen["data"]
    assert "email=" not in body and "api_key=" not in body
    assert "MEGABLAST=on" in body
    assert "DATABASE=core_nt" in body
    assert "EXPECT=0.05" in body
    assert "WORD_SIZE=28" in body
    assert "NUCL_REWARD=1" in body
    assert "NUCL_PENALTY=-2" in body
    assert "FILTER=L" in body and "FILTER=m" in body
    assert "GAPCOSTS=0+0" in body or "GAPCOSTS=0%200" in body or "GAPCOSTS=0 0" in body

    import warnings
    from Bio import Entrez

    Entrez.email = "leftover@example.com"
    Entrez.api_key = "leftover-key"
    ncbi_clear = NCBI(log=lambda *_: None)
    assert ncbi_clear.Entrez.email is None
    assert ncbi_clear.Entrez.api_key is None
    with warnings.catch_warnings(record=True) as caught:
        Entrez._construct_params({})
    assert not any("mail" in str(w.message).lower() for w in caught)

    monkeypatch.setenv("NCBI_EMAIL", "person@example.com")
    monkeypatch.setenv("NCBI_API_KEY", "secret-key")
    params = blast_identity_params()
    assert ("email", "person@example.com") in params
    assert all(k != "api_key" for k, _ in params)
    ncbi2 = NCBI(log=lambda *_: None)
    assert ncbi2.Entrez.email == "person@example.com"
    assert ncbi2.Entrez.api_key == "secret-key"


def test_reuse_rids_skips_submit():
    seq = "ACGT" * 20
    fake = FakeNCBI([("S", seq, [_spec("Pyxine copelandii", "EU670225.1", len(seq), len(seq), len(seq))])])

    def nope(_fasta):
        raise AssertionError("submit 되면 안 됩니다")

    fake.submit = nope
    queries = [dict(key=0, label="S", seq=seq)]
    res = run_blast(fake, queries, 10, "/tmp", log=lambda *_: None, reuse=["FAKERID"])
    assert res[0]["hits"][0]["accession"] == "EU670225.1"


def test_template_roundtrip_values_equal_web_values(tmp_path):
    """헤더만 있는 시트에 필드 값을 쓰고 다시 읽어 web_values 와 같다."""
    seq = "A" * 320
    spec = {
        "name": "Lecidella biformis",
        "taxid": 1,
        "acc": "PZ322662.1",
        "length": 540,
        "hsps": [
            {"bit": 100.2, "ident": 50, "alen": 60, "qfrom": 1, "qto": 60, "evalue": 1e-5},
            {"bit": 400.2, "ident": 200, "alen": 210, "qfrom": 100, "qto": 300, "evalue": 1.23e-50},
        ],
    }
    v = web_values(json_hit(spec), len(seq))
    tmpl = load_template()
    wb = openpyxl.Workbook()
    ws = wb.active
    # 결과 헤더를 역순으로 배치
    headers = ["Accession", "Per. Ident", "Query Cover", "E value", "Total Score", "Max Score", "Scientific Name"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    cols = map_fields(ws, tmpl)
    cols = ensure_fields(ws, tmpl, cols, log=lambda *_: None)
    blast = {
        "hits": [v],
        "lineage": "Fungi; Dikarya",
        "rid": "R",
        "rid_url": "https://blast.ncbi.nlm.nih.gov/Blast.cgi?CMD=Get&RID=R",
        "qlen": len(seq),
    }
    values = build_field_values(sample_id="S", sequence=seq, blast=blast, level=None, comment="", merge=None)
    write_fields(ws, 2, tmpl, cols, values, fill_input=True)
    style_sheet(ws, tmpl, cols, {})
    assert ws.cell(2, cols["scientific_name"]).value == v["name"]
    assert ws.cell(2, cols["max_score"]).value == 400
    assert ws.cell(2, cols["total_score"]).value == 500
    assert ws.cell(2, cols["query_cover"]).value == 0.82
    assert ws.cell(2, cols["per_ident"]).value == 0.9524
    assert ws.cell(2, cols["evalue"]).value == pytest.approx(1e-50)
    assert ws.cell(2, cols["accession"]).value == "PZ322662.1"
    # 독립 파서
    xml = blast_xml([("S", seq, [spec])])
    from blastid.scoring import parse_xml_report

    parsed = parse_xml_report(xml)[0]
    assert parsed["pi"] == v["pi"]
    assert parsed["qc"] == v["qc"]
    assert parsed["max"] == v["max"]


def _args(**kw):
    base = dict(
        kind="excel",
        inputs=[],
        output=None,
        out_dir=None,
        template=None,
        sheet=None,
        rows=None,
        bold_only=False,
        only_empty=False,
        in_place=False,
        clear_old_comments=False,
        dry_run=False,
        batch_size=10,
        min_ident=90.0,
        min_cover=80.0,
        alt_margin=2.0,
        no_verify=False,
        reuse_rids=None,
        min_seq_len=50,
        id_header=None,
        seq_header=None,
        blast=True,
        pair_regex=None,
        pairs=None,
        no_merge=False,
        merger="biopython",
        tracy_trim=4,
        trimmer="mott",
        trim_cutoff=0.05,
        min_len=150,
        min_q=25,
        min_overlap=50,
        min_overlap_ident=90.0,
        conflict_q=5,
        indel_q=15,
    )
    base.update(kw)

    class A:
        pass

    a = A()
    for k, v in base.items():
        setattr(a, k, v)
    return a
