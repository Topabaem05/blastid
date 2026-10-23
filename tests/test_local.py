"""24시간 한도를 넘기면 로컬 BLAST+로 넘긴다. IP 우회는 하지 않는다."""
import random
from pathlib import Path

import openpyxl
import pytest

from blastid.fields import build_field_values
from blastid.localblast import LocalBlastUnavailable, make_local_db, resolve_db
from blastid.pipeline import run_blast
from blastid.policy import UsageLog
from blastid.template import load_template
from blastid.workbook import style_sheet, write_fields
from tests.blastlib import blast_json

BLASTN = Path("/tmp/ncbi-blast/ncbi-blast-2.17.0+/bin/blastn")


def _seq() -> str:
    random.seed(1)
    return "".join(random.choice("ACGT") for _ in range(180))


def _spec():
    return {
        "name": "Remote species",
        "taxid": 42,
        "acc": "EU000001.1",
        "title": "Remote species",
        "length": 180,
        "hsps": [{"bit": 300.2, "ident": 180, "alen": 180, "qfrom": 1, "qto": 180, "evalue": 1e-80}],
    }


class _Fake:
    def __init__(self, queries):
        self.email = ""
        self.fast = True
        self._last_poll = {}
        self.calls = 0
        self.json = blast_json(queries)
        self.tax = {}

    def submit(self, _fasta):
        self.calls += 1
        return "RIDREMOTE", 0

    def wait_all(self, jobs, max_wait=3600):
        import time

        return {rid: time.time() for rid in jobs}

    def get_json(self, _rid):
        return self.json

    def taxonomy(self, _ids):
        return {}


def test_switch_keeps_remote_rows_when_local_blast_is_missing(tmp_path):
    seq = "ACGT" * 30
    fake = _Fake([("Keep", seq, [_spec()])])
    queries = [
        dict(key="a", label="Keep", seq=seq),
        dict(key="b", label="Next", seq=seq),
    ]
    with pytest.raises(LocalBlastUnavailable) as caught:
        run_blast(
            fake,
            queries,
            1,
            tmp_path,
            log=lambda *_: None,
            usage=UsageLog(tmp_path / "usage.json"),
            backend="auto",
            remote_limit=1,
            local_db="missing-db",
            blastn=str(tmp_path / "no-such-blastn"),
        )
    assert fake.calls == 1
    assert "a" in caught.value.partial
    assert "b" not in caught.value.partial
    assert "setup-local" in str(caught.value)
    assert "only-empty" in str(caught.value)


@pytest.mark.skipif(not BLASTN.is_file(), reason="NCBI BLAST+ 2.17 binary is not in /tmp/ncbi-blast")
def test_over_limit_uses_local_database(tmp_path):
    seq = _seq()
    db = tmp_path / "tiny"
    fasta = tmp_path / "sub.fa"
    fasta.write_text(f">SYN1.1 Pyxine testensis ITS\n{seq}\n", encoding="utf-8")
    make_local_db(fasta, db, str(BLASTN))
    assert resolve_db(str(db)) == db
    usage = UsageLog(tmp_path / "usage.json")
    for i in range(100):
        usage.add(rid=f"old{i}")
    fake = _Fake([])

    def _no_submit(_fasta):
        raise AssertionError("한도를 넘겼는데 원격 제출이 일어났습니다")

    fake.submit = _no_submit
    queries = [dict(key="local", label="q1", seq=seq)]
    results = run_blast(
        fake,
        queries,
        10,
        tmp_path / "out",
        log=lambda *_: None,
        usage=usage,
        backend="auto",
        remote_limit=100,
        local_db=str(db),
        blastn=str(BLASTN),
    )
    hit = results["local"]
    assert hit["backend"] == "local"
    assert hit["source"].startswith("로컬 BLAST+ (")
    assert hit["hits"][0]["name"] == "Pyxine testensis"
    assert hit["hits"][0]["accession"] == "SYN1.1"
    assert hit["hits"][0]["pi"] == 1
    assert hit["hits"][0]["qc"] == 1
    assert usage.count() == 100


def test_template_marks_local_source(tmp_path):
    tmpl = load_template()
    assert any(col.field == "source" and col.header == "검색 출처" for col in tmpl.columns)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tmpl.sheet
    for index, col in enumerate(tmpl.columns, start=1):
        ws.cell(1, index, col.header)
    from blastid.workbook import map_fields

    cols = map_fields(ws, tmpl)
    blast = {
        "hits": [{
            "name": "Pyxine testensis",
            "max": 300,
            "total": 300,
            "qc": 1,
            "evalue": 0,
            "pi": 1,
            "acc_len": 180,
            "accession": "SYN1.1",
            "taxid": None,
            "max_s": "300",
            "total_s": "300",
            "qc_s": "100%",
            "e_s": "0.0",
            "pi_s": "100.00%",
        }],
        "lineage": "",
        "rid": "local:tiny",
        "rid_url": "",
        "qlen": 180,
        "source": "로컬 BLAST+ (tiny)",
        "backend": "local",
    }
    values = build_field_values(sample_id="q1", sequence="A" * 180, blast=blast, level=None, comment="로컬", merge=None)
    write_fields(ws, 2, tmpl, cols, values, fill_input=True)
    style_sheet(ws, tmpl, cols, {}, local_rows={2})
    source_col = cols["source"]
    assert ws.cell(2, source_col).value == "로컬 BLAST+ (tiny)"
    assert ws.cell(2, source_col).fill.fgColor.rgb.endswith("D6E3F0")
    assert ws.cell(2, cols["scientific_name"]).fill.fgColor.rgb.endswith("D6E3F0")
    wb.save(tmp_path / "marked.xlsx")
